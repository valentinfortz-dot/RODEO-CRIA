from flask import Flask, render_template, request, jsonify, send_file
from database import init_db, get_db
import os, io, csv, tempfile

app = Flask(__name__)
app.secret_key = os.environ.get('SECRET_KEY', 'rodeo-secret-2024')

with app.app_context():
    init_db()
    try:
        from seed_render import seed_if_empty
        seed_if_empty()
    except Exception as e:
        print(f"Seed omitido: {e}")

@app.before_request
def setup():
    init_db()

# ─── INICIO ───────────────────────────────────────────────
@app.route('/')
def index():
    db = get_db()
    stats = {
        'vacas':    db.execute("SELECT COUNT(*) FROM vacas WHERE estado='Activa'").fetchone()[0],
        'toros':    db.execute("SELECT COUNT(*) FROM toros").fetchone()[0],
        'terneros': db.execute("SELECT COUNT(*) FROM terneros").fetchone()[0],
        'vendidas': db.execute("SELECT COUNT(*) FROM vacas WHERE estado!='Activa'").fetchone()[0],
    }
    temporadas = [r['temporada'] for r in db.execute("SELECT DISTINCT temporada FROM eventos ORDER BY temporada DESC").fetchall()]
    db.close()
    return render_template('index.html', stats=stats, temporadas=temporadas)

# ─── BUSCAR ───────────────────────────────────────────────
@app.route('/buscar')
def buscar():
    q = request.args.get('q','').strip()
    if not q:
        return render_template('buscar.html', resultados=[], q='')
    db = get_db()
    resultados = db.execute(
        "SELECT * FROM vacas WHERE ide LIKE ? ORDER BY estado, ide LIMIT 30",
        (f'%{q.replace(" ","")}%',)
    ).fetchall()
    db.close()
    return render_template('buscar.html', resultados=resultados, q=q)

# ─── FICHA VACA ───────────────────────────────────────────
@app.route('/vaca/<ide>')
def ficha_vaca(ide):
    db = get_db()
    vaca = db.execute("SELECT * FROM vacas WHERE ide=?", (ide,)).fetchone()
    if not vaca:
        db.close()
        return render_template('error.html', msg=f'Vaca {ide} no encontrada.'), 404
    eventos  = db.execute("SELECT * FROM eventos WHERE ide_vaca=? ORDER BY temporada DESC", (ide,)).fetchall()
    terneros = db.execute("SELECT * FROM terneros WHERE ide_madre=? ORDER BY temporada DESC", (ide,)).fetchall()
    madre = db.execute("SELECT * FROM vacas WHERE ide=?", (vaca['ide_madre'],)).fetchone() if vaca['ide_madre'] else None
    hijos = db.execute("SELECT * FROM vacas WHERE ide_madre=?", (ide,)).fetchall()
    db.close()
    return render_template('ficha_vaca.html', vaca=vaca, eventos=eventos, terneros=terneros, madre=madre, hijos=hijos)

# ─── NACIMIENTO ───────────────────────────────────────────
@app.route('/nacimiento', methods=['GET','POST'])
def nacimiento():
    if request.method == 'POST':
        data = request.form
        ide_madre = data.get('ide_madre','').strip().replace(' ','')
        caravana  = data.get('caravana','').strip()
        temporada = int(data.get('temporada', 2025))
        db = get_db()
        if not db.execute("SELECT ide FROM vacas WHERE ide=?", (ide_madre,)).fetchone():
            db.close()
            return jsonify({'ok': False, 'error': f'Vaca {ide_madre} no encontrada.'})
        if db.execute("SELECT id FROM terneros WHERE caravana=? AND temporada=?", (caravana, temporada)).fetchone():
            db.close()
            return jsonify({'ok': False, 'error': f'Caravana {caravana} ya registrada en {temporada}.'})
        db.execute('INSERT INTO terneros (caravana,ide_madre,temporada,sexo,raza,semana_nacimiento,potrero) VALUES (?,?,?,?,?,?,?)',
                   (caravana, ide_madre, temporada, data.get('sexo'), data.get('raza'), data.get('semana'), data.get('potrero')))
        db.execute("UPDATE eventos SET caravana_ternero=? WHERE ide_vaca=? AND temporada=? AND (caravana_ternero IS NULL OR caravana_ternero='')",
                   (caravana, ide_madre, temporada))
        db.commit(); db.close()
        return jsonify({'ok': True, 'msg': f'✓ Ternero {caravana} registrado. Madre: {ide_madre}'})
    return render_template('nacimiento.html', potreros=['T','N','M','F1','F2','R','LP'], razas=['AA','RA','HA','HE'])

# ─── NUEVO EVENTO ─────────────────────────────────────────
@app.route('/evento/nuevo', methods=['GET','POST'])
def nuevo_evento():
    db = get_db()
    if request.method == 'POST':
        data = request.form
        ide = data.get('ide_vaca','').strip().replace(' ','')
        temporada = int(data.get('temporada', 2025))
        if not db.execute("SELECT ide FROM vacas WHERE ide=?", (ide,)).fetchone():
            db.close()
            return jsonify({'ok': False, 'error': f'Vaca {ide} no encontrada.'})
        campos = (data.get('dao_sincro') or None, data.get('toro_iatf') or None,
                  data.get('diagnostico_iatf') or None, data.get('ecografia_final') or None,
                  data.get('toro_entore') or None, data.get('eco_entore') or None,
                  data.get('observaciones') or None)
        if db.execute("SELECT id FROM eventos WHERE ide_vaca=? AND temporada=?", (ide, temporada)).fetchone():
            db.execute('UPDATE eventos SET dao_sincro=?,toro_iatf=?,diagnostico_iatf=?,ecografia_final=?,toro_entore=?,eco_entore=?,observaciones=? WHERE ide_vaca=? AND temporada=?',
                       campos + (ide, temporada))
        else:
            db.execute('INSERT INTO eventos (ide_vaca,temporada,dao_sincro,toro_iatf,diagnostico_iatf,ecografia_final,toro_entore,eco_entore,observaciones) VALUES (?,?,?,?,?,?,?,?,?)',
                       (ide, temporada) + campos)
        if data.get('eco_entore') == 'VACIA':
            db.execute("UPDATE vacas SET estado='Vendida', motivo_baja=? WHERE ide=?", (f'Vacía {temporada}', ide))
        db.commit(); db.close()
        return jsonify({'ok': True, 'msg': f'✓ Evento {temporada} guardado para vaca {ide}'})
    toros = db.execute("SELECT * FROM toros ORDER BY idv").fetchall()
    ide_prefill = request.args.get('ide','')
    db.close()
    return render_template('nuevo_evento.html', toros=toros, ide_prefill=ide_prefill)

# ─── NUEVA VACA ───────────────────────────────────────────
@app.route('/vaca/nueva', methods=['GET','POST'])
def nueva_vaca():
    db = get_db()
    if request.method == 'POST':
        data = request.form
        ide = data.get('ide','').strip().replace(' ','')
        if not ide:
            db.close()
            return jsonify({'ok': False, 'error': 'IDE requerido.'})
        if db.execute("SELECT id FROM vacas WHERE ide=?", (ide,)).fetchone():
            db.close()
            return jsonify({'ok': False, 'error': f'Vaca {ide} ya existe.'})
        db.execute('INSERT INTO vacas (ide,anio_nacimiento,origen,clase,estado,ide_madre) VALUES (?,?,?,?,?,?)',
                   (ide, data.get('anio_nacimiento') or None, data.get('origen') or None,
                    data.get('clase') or None, 'Activa', data.get('ide_madre') or None))
        db.commit(); db.close()
        return jsonify({'ok': True, 'msg': f'✓ Vaca {ide} registrada.', 'ide': ide})
    db.close()
    return render_template('nueva_vaca.html')

# ─── DASHBOARD ────────────────────────────────────────────
@app.route('/dashboard')
def dashboard():
    db = get_db()
    temporada = int(request.args.get('temporada', 2025))
    temporadas = [r['temporada'] for r in db.execute("SELECT DISTINCT temporada FROM eventos ORDER BY temporada DESC").fetchall()]

    # % preñez general por temporada
    prenez_general = db.execute('''
        SELECT temporada,
               COUNT(*) as total,
               SUM(CASE WHEN eco_entore IN ('GRANDE','CHICA','PREÑADA') OR diagnostico_iatf IN ('PRENA','SUPERFICIAL','PROFUNDO') THEN 1 ELSE 0 END) as prenadas
        FROM eventos GROUP BY temporada ORDER BY temporada DESC
    ''').fetchall()

    # % preñez por toro IATF
    toros_iatf = db.execute('''
        SELECT toro_iatf as toro, COUNT(*) as total,
               SUM(CASE WHEN diagnostico_iatf IN ('PRENA','SUPERFICIAL','PROFUNDO') THEN 1 ELSE 0 END) as prenadas
        FROM eventos WHERE temporada=? AND toro_iatf IS NOT NULL AND toro_iatf NOT IN ('ENTORE','')
        GROUP BY toro_iatf ORDER BY prenadas*1.0/total DESC
    ''', (temporada,)).fetchall()

    # % preñez por toro entore
    toros_entore = db.execute('''
        SELECT toro_entore as toro, COUNT(*) as total,
               SUM(CASE WHEN eco_entore IN ('GRANDE','CHICA','PREÑADA') THEN 1 ELSE 0 END) as prenadas
        FROM eventos WHERE temporada=? AND toro_entore IS NOT NULL AND toro_entore NOT IN ('ENTORE','-','')
        GROUP BY toro_entore ORDER BY prenadas*1.0/total DESC
    ''', (temporada,)).fetchall()

    # % preñez por edad de la vaca
    prenez_edad = db.execute('''
        SELECT v.anio_nacimiento, COUNT(*) as total,
               SUM(CASE WHEN e.eco_entore IN ('GRANDE','CHICA','PREÑADA') OR e.diagnostico_iatf IN ('PRENA','SUPERFICIAL','PROFUNDO') THEN 1 ELSE 0 END) as prenadas
        FROM eventos e JOIN vacas v ON e.ide_vaca=v.ide
        WHERE e.temporada=? AND v.anio_nacimiento IS NOT NULL
        GROUP BY v.anio_nacimiento ORDER BY v.anio_nacimiento
    ''', (temporada,)).fetchall()

    # Peso destete por sexo
    peso_sexo = db.execute('''
        SELECT sexo, ROUND(AVG(peso_destete),1) as prom, COUNT(*) as n,
               ROUND(MIN(peso_destete),1) as min, ROUND(MAX(peso_destete),1) as max
        FROM terneros WHERE temporada=? AND peso_destete IS NOT NULL AND peso_destete > 0 AND sexo IS NOT NULL
        GROUP BY sexo
    ''', (temporada,)).fetchall()

    # Peso destete por padre (toro entore)
    peso_padre = db.execute('''
        SELECT e.toro_entore as padre,
               ROUND(AVG(t.peso_destete),1) as prom_peso,
               ROUND(AVG(t.gmd),3) as prom_gmd,
               COUNT(t.id) as n_terneros
        FROM terneros t
        JOIN eventos e ON t.ide_madre=e.ide_vaca AND t.temporada=e.temporada
        WHERE t.temporada=? AND t.peso_destete IS NOT NULL AND t.peso_destete > 0
              AND e.toro_entore IS NOT NULL AND e.toro_entore NOT IN ('ENTORE','-','')
        GROUP BY e.toro_entore ORDER BY prom_peso DESC
    ''', (temporada,)).fetchall()

    # Peso destete por edad de la madre
    peso_edad_madre = db.execute('''
        SELECT v.anio_nacimiento as anio_madre,
               ROUND(AVG(t.peso_destete),1) as prom_peso,
               ROUND(AVG(t.gmd),3) as prom_gmd,
               COUNT(t.id) as n
        FROM terneros t
        JOIN vacas v ON t.ide_madre=v.ide
        WHERE t.temporada=? AND t.peso_destete IS NOT NULL AND t.peso_destete > 0
              AND v.anio_nacimiento IS NOT NULL
        GROUP BY v.anio_nacimiento ORDER BY v.anio_nacimiento
    ''', (temporada,)).fetchall()

    # GMD por temporada
    gmd_temporada = db.execute('''
        SELECT temporada, ROUND(AVG(gmd),3) as gmd_prom,
               ROUND(AVG(CASE WHEN sexo='M' THEN gmd END),3) as gmd_m,
               ROUND(AVG(CASE WHEN sexo='H' THEN gmd END),3) as gmd_h,
               COUNT(*) as n
        FROM terneros WHERE gmd IS NOT NULL AND gmd > 0
        GROUP BY temporada ORDER BY temporada
    ''').fetchall()

    # Candidatas a baja
    candidatas = db.execute('''
        SELECT v.ide, v.anio_nacimiento, v.origen,
               GROUP_CONCAT(e.temporada||':'||COALESCE(e.eco_entore,e.diagnostico_iatf,'?'), ' | ') as resumen
        FROM vacas v JOIN eventos e ON v.ide=e.ide_vaca
        WHERE v.estado='Activa'
        GROUP BY v.ide
        HAVING SUM(CASE WHEN e.eco_entore IN ('VACIA') OR e.eco_entore IS NULL THEN 1 ELSE 0 END) >= 2
        ORDER BY v.ide LIMIT 20
    ''').fetchall()

    db.close()
    return render_template('dashboard.html',
        temporada=temporada, temporadas=temporadas,
        prenez_general=prenez_general,
        toros_iatf=toros_iatf, toros_entore=toros_entore,
        prenez_edad=prenez_edad,
        peso_sexo=peso_sexo, peso_padre=peso_padre,
        peso_edad_madre=peso_edad_madre,
        gmd_temporada=gmd_temporada,
        candidatas=candidatas)

# ─── API VACA ─────────────────────────────────────────────
@app.route('/api/vaca/<ide>')
def api_vaca(ide):
    db = get_db()
    vaca = db.execute("SELECT ide,anio_nacimiento,clase,estado FROM vacas WHERE ide=?", (ide.strip().replace(' ',''),)).fetchone()
    db.close()
    if vaca:
        return jsonify({'encontrada': True, 'ide': vaca['ide'], 'anio': vaca['anio_nacimiento'],
                        'clase': vaca['clase'], 'estado': vaca['estado']})
    return jsonify({'encontrada': False})

# ─── IMPORTACIÓN MASIVA ───────────────────────────────────
@app.route('/importar')
def importar():
    return render_template('importar.html')

def _importar_csv(file, temporada, fn):
    content = file.read().decode('utf-8-sig')
    reader = csv.DictReader(io.StringIO(content))
    db = get_db(); ok = 0; errores = []
    for row in reader:
        result = fn(db, row, temporada)
        if result is True: ok += 1
        elif isinstance(result, str): errores.append(result)
    db.commit(); db.close()
    return ok, errores

@app.route('/importar/eco', methods=['POST'])
def importar_eco():
    file = request.files.get('archivo')
    temporada = int(request.form.get('temporada', 2025))
    campo = request.form.get('campo', 'eco_entore')
    if not file: return jsonify({'ok': False, 'error': 'No se subió ningún archivo.'})
    def fn(db, row, temp):
        ide = str(row.get('IDE','')).strip().replace(' ','')
        val = str(row.get('VALOR', row.get('ECO ENTORE', row.get('DIAGNOSTICO', row.get('RESULTADO',''))))).strip().upper()
        if not ide or not val or val in ['NAN','']: return None
        if not db.execute("SELECT ide FROM vacas WHERE ide=?", (ide,)).fetchone(): return f'{ide} no encontrada'
        if db.execute("SELECT id FROM eventos WHERE ide_vaca=? AND temporada=?", (ide, temp)).fetchone():
            db.execute(f"UPDATE eventos SET {campo}=? WHERE ide_vaca=? AND temporada=?", (val, ide, temp))
        else:
            db.execute(f"INSERT INTO eventos (ide_vaca,temporada,{campo}) VALUES (?,?,?)", (ide, temp, val))
        if campo == 'eco_entore' and val == 'VACIA':
            db.execute("UPDATE vacas SET estado='Vendida', motivo_baja=? WHERE ide=?", (f'Vacía {temp}', ide))
        return True
    ok, errores = _importar_csv(file, temporada, fn)
    return jsonify({'ok': True, 'msg': f'✓ {ok} registros actualizados.', 'errores': errores[:5]})

@app.route('/importar/nacimientos', methods=['POST'])
def importar_nacimientos():
    file = request.files.get('archivo')
    temporada = int(request.form.get('temporada', 2025))
    if not file: return jsonify({'ok': False, 'error': 'No se subió ningún archivo.'})
    def fn(db, row, temp):
        ide_m = str(row.get('IDE MADRE', row.get('IDE',''))).strip().replace(' ','')
        car   = str(row.get('CARAVANA', row.get('RP',''))).strip()
        if not ide_m or not car: return None
        if not db.execute("SELECT ide FROM vacas WHERE ide=?", (ide_m,)).fetchone(): return f'{ide_m} no encontrada'
        if db.execute("SELECT id FROM terneros WHERE caravana=? AND temporada=?", (car, temp)).fetchone(): return None
        try: peso = float(row['PESO DESTETE']) if row.get('PESO DESTETE') else None
        except: peso = None
        try: dias = int(float(row['DIAS DESTETE'])) if row.get('DIAS DESTETE') else None
        except: dias = None
        try: gmd = round(float(row['GMD']),3) if row.get('GMD') else None
        except: gmd = None
        db.execute('INSERT INTO terneros (caravana,ide_madre,temporada,sexo,raza,semana_nacimiento,potrero,peso_destete,dias_destete,gmd) VALUES (?,?,?,?,?,?,?,?,?,?)',
                   (car,ide_m,temp,
                    str(row.get('SEXO','')).strip() or None,
                    str(row.get('RAZA','')).strip() or None,
                    str(row.get('SEMANA','')).strip() or None,
                    str(row.get('POTRERO','')).strip() or None,
                    peso, dias, gmd))
        return True
    ok, errores = _importar_csv(file, temporada, fn)
    return jsonify({'ok': True, 'msg': f'✓ {ok} terneros registrados.', 'errores': errores[:5]})

@app.route('/importar/iatf', methods=['POST'])
def importar_iatf():
    file = request.files.get('archivo')
    temporada = int(request.form.get('temporada', 2025))
    if not file: return jsonify({'ok': False, 'error': 'No se subió ningún archivo.'})
    def fn(db, row, temp):
        ide   = str(row.get('IDE','')).strip().replace(' ','')
        dao   = str(row.get('DAO SINCRO','')).strip().upper() or None
        tiatf = str(row.get('TORO IATF','')).strip().upper() or None
        tent  = str(row.get('TORO ENTORE','')).strip().upper() or None
        if not ide: return None
        if not db.execute("SELECT ide FROM vacas WHERE ide=?", (ide,)).fetchone(): return f'{ide} no encontrada'
        if db.execute("SELECT id FROM eventos WHERE ide_vaca=? AND temporada=?", (ide, temp)).fetchone():
            if dao:   db.execute("UPDATE eventos SET dao_sincro=? WHERE ide_vaca=? AND temporada=?",  (dao,ide,temp))
            if tiatf: db.execute("UPDATE eventos SET toro_iatf=? WHERE ide_vaca=? AND temporada=?",  (tiatf,ide,temp))
            if tent:  db.execute("UPDATE eventos SET toro_entore=? WHERE ide_vaca=? AND temporada=?",(tent,ide,temp))
        else:
            db.execute('INSERT INTO eventos (ide_vaca,temporada,dao_sincro,toro_iatf,toro_entore) VALUES (?,?,?,?,?)',
                       (ide,temp,dao,tiatf,tent))
        return True
    ok, errores = _importar_csv(file, temporada, fn)
    return jsonify({'ok': True, 'msg': f'✓ {ok} registros IATF actualizados.', 'errores': errores[:5]})

# ─── PLANTILLAS DESCARGABLES ──────────────────────────────
@app.route('/plantilla/<tipo>')
def plantilla(tipo):
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError:
        return "openpyxl no instalado", 500

    db = get_db()
    vacas = db.execute("SELECT ide, anio_nacimiento, clase, origen FROM vacas WHERE estado='Activa' ORDER BY ide").fetchall()
    db.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    green   = PatternFill('solid', fgColor='0F6E56')
    amber   = PatternFill('solid', fgColor='BA7517')
    amber_l = PatternFill('solid', fgColor='FAEEDA')
    gray_l  = PatternFill('solid', fgColor='F1EFE8')
    white   = PatternFill('solid', fgColor='FFFFFF')

    configs = {
        'eco':         ('ECO ENTORE',         ['IDE','Año','Clase','Origen','VALOR'],               '"GRANDE,CHICA,VACIA,ABORTO,VENTA"', 4, 'eco_entore'),
        'diagnostico': ('DIAGNOSTICO IATF',   ['IDE','Año','Clase','Origen','VALOR'],               '"PRENA,CICLA,SUPERFICIAL,PROFUNDO,ENTORE"', 4, 'diagnostico_iatf'),
        'iatf':        ('INSEMINACION IATF',  ['IDE','Año','Clase','Origen','TORO IATF','DAO SINCRO','TORO ENTORE'], None, 4, None),
        'nacimientos': ('NACIMIENTOS',        ['IDE MADRE','Año','Clase','CARAVANA','SEXO','RAZA','POTRERO','SEMANA','PESO DESTETE','DIAS DESTETE','GMD'], None, None, None),
    }
    title, cols, validacion, col_key, _ = configs.get(tipo, configs['eco'])
    ws.title = title

    ws.merge_cells(f'A1:{chr(64+len(cols))}1')
    ws['A1'].value = f'📋  {title} — completá y guardá como CSV para importar'
    ws['A1'].font  = Font(bold=True, color='FFFFFF', size=12)
    ws['A1'].fill  = green
    ws['A1'].alignment = Alignment(horizontal='left', vertical='center', indent=1)
    ws.row_dimensions[1].height = 24

    for ci, h in enumerate(cols, 1):
        cell = ws.cell(row=2, column=ci)
        cell.value = h
        cell.font  = Font(bold=True, color='FFFFFF', size=11)
        cell.fill  = amber if (col_key and ci > col_key) else green
        cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[2].height = 22

    dv = None
    if validacion and col_key:
        dv = DataValidation(type="list", formula1=validacion, allow_blank=True, showDropDown=False)
        ws.add_data_validation(dv)

    for i, vaca in enumerate(vacas):
        r = i + 3
        fill = gray_l if i % 2 == 0 else white
        base = [vaca['ide'], vaca['anio_nacimiento'] or '', vaca['clase'] or '', vaca['origen'] or '']
        extras = [''] * (len(cols) - 4)
        row_vals = (base + extras)[:len(cols)]
        for ci, val in enumerate(row_vals, 1):
            cell = ws.cell(row=r, column=ci)
            cell.value = val
            cell.fill  = amber_l if (col_key and ci > col_key) else fill
            cell.font  = Font(size=10, color='888888' if (col_key and ci <= col_key) else '000000')
        if dv and col_key:
            dv.add(f'{chr(64+col_key+1)}{r}')

    for ci in range(1, len(cols)+1):
        ws.column_dimensions[chr(64+ci)].width = 22
    ws.freeze_panes = 'A3'

    tmp = tempfile.NamedTemporaryFile(delete=False, suffix='.xlsx')
    wb.save(tmp.name); tmp.close()
    return send_file(tmp.name, as_attachment=True, download_name=f'{title.replace(" ","_")}.xlsx')

if __name__ == '__main__':
    port = int(os.environ.get('PORT', 5000))
    app.run(host='0.0.0.0', port=port, debug=False)
